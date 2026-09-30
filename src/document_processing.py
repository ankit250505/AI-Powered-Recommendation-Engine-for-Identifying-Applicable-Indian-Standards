from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path


SUPPORTED_EXTENSIONS = {".txt", ".pdf", ".docx"}
MAX_UPLOAD_BYTES = 10 * 1024 * 1024


class DocumentProcessingError(ValueError):
    pass


@dataclass(frozen=True)
class ExtractedDocument:
    filename: str
    extension: str
    text: str

    @property
    def char_count(self) -> int:
        return len(self.text)

    @property
    def preview(self) -> str:
        compact = " ".join(self.text.split())
        return compact[:800]


def extract_text_from_upload(filename: str, content: bytes) -> ExtractedDocument:
    extension = Path(filename).suffix.lower()

    if len(content) > MAX_UPLOAD_BYTES:
        raise DocumentProcessingError("File is too large. Maximum supported upload size is 10 MB.")

    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise DocumentProcessingError(f"Unsupported file type. Supported formats: {supported}.")

    if extension == ".txt":
        text = extract_txt(content)
    elif extension == ".pdf":
        text = extract_pdf(content)
    else:
        text = extract_docx(content)

    text = clean_text(text)
    if not text:
        if extension == ".pdf":
            raise DocumentProcessingError(
                "No selectable text was found. This may be a scanned PDF; run OCR to create "
                "a searchable PDF, then upload it again."
            )
        raise DocumentProcessingError("No readable text could be extracted from this document.")

    return ExtractedDocument(filename=filename, extension=extension, text=text)


def extract_txt(content: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "cp1252"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="ignore")


def extract_pdf(content: bytes) -> str:
    try:
        import fitz
    except ModuleNotFoundError as error:
        raise DocumentProcessingError("PDF extraction dependency is not installed.") from error

    try:
        document = fitz.open(stream=content, filetype="pdf")
    except Exception as error:
        raise DocumentProcessingError("Could not open PDF file.") from error

    parts: list[str] = []
    with document:
        for page in document:
            parts.append(page.get_text("text"))
    return "\n".join(parts)


def extract_docx(content: bytes) -> str:
    try:
        from docx import Document
    except ModuleNotFoundError as error:
        raise DocumentProcessingError("DOCX extraction dependency is not installed.") from error

    try:
        document = Document(BytesIO(content))
    except Exception as error:
        raise DocumentProcessingError("Could not open DOCX file.") from error

    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]

    table_text: list[str] = []
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                table_text.append(" | ".join(cells))

    return "\n".join(paragraphs + table_text)


def clean_text(text: str) -> str:
    lines = [" ".join(line.split()) for line in text.splitlines()]
    return "\n".join(line for line in lines if line)
